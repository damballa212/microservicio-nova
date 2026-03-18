import io
import os
import tempfile
from typing import Any, cast

import httpx
from langchain_text_splitters import RecursiveCharacterTextSplitter


def _read_bytes_from_source(path: str | None, presigned_url: str | None) -> bytes:
    if isinstance(path, str) and os.path.exists(path):
        return open(path, "rb").read()
    if isinstance(presigned_url, str) and presigned_url:
        with httpx.Client(timeout=30.0) as client:
            r = client.get(presigned_url)
            r.raise_for_status()
            return r.content
    return b""


def _detect_kind(mime: str | None, filename: str | None) -> str:
    m = (mime or "").lower().strip()
    f = (filename or "").lower().strip()
    if m in {
        "application/pdf",
    } or f.endswith(".pdf"):
        return "pdf"
    if m in {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    } or f.endswith(".docx"):
        return "docx"
    if m in {
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    } or f.endswith(".xlsx"):
        return "xlsx"
    if m in {"text/csv"} or f.endswith(".csv"):
        return "csv"
    if m.startswith("text/") or f.endswith(".txt"):
        return "text"
    return "binary"


def _extract_text_pdf(data: bytes) -> str:
    try:
        from pdfminer.high_level import extract_text

        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf")
        try:
            tmp.write(data)
            tmp.flush()
            return extract_text(tmp.name) or ""
        finally:
            try:
                tmp.close()
                os.unlink(tmp.name)
            except Exception:
                pass
    except Exception:
        return ""


def _extract_text_docx(data: bytes) -> str:
    try:
        import docx

        bio = io.BytesIO(data)
        doc = docx.Document(bio)
        parts = [p.text for p in doc.paragraphs if isinstance(p.text, str)]
        return "\n".join([t.strip() for t in parts if t and t.strip()])
    except Exception:
        return ""


def _extract_rows_xlsx(data: bytes) -> list[dict[str, Any]]:
    try:
        import openpyxl

        bio = io.BytesIO(data)
        wb = openpyxl.load_workbook(bio, data_only=True, read_only=True)
        ws = wb.active
        rows = []
        header = None
        for i, r in enumerate(ws.iter_rows(values_only=True)):
            vals = [(v if v is not None else "") for v in r]
            if i == 0:
                header = [str(v) for v in vals]
                continue
            if not header:
                continue
            obj = {}
            for j, h in enumerate(header):
                obj[str(h)] = str(vals[j]) if j < len(vals) else ""
            rows.append(obj)
        return rows
    except Exception:
        return []


def _extract_rows_csv(data: bytes) -> list[dict[str, Any]]:
    try:
        import csv

        import chardet

        enc = chardet.detect(data or b""
        ).get("encoding") or "utf-8"
        s = data.decode(enc, errors="ignore")
        bio = io.StringIO(s)
        reader = csv.reader(bio)
        rows_raw = [list(r) for r in reader]
        if not rows_raw:
            return []
        header = [str(v or "") for v in rows_raw[0]]
        out = []
        for r in rows_raw[1:]:
            obj = {}
            for j, h in enumerate(header):
                obj[str(h)] = str(r[j] if j < len(r) else "")
            out.append(obj)
        return out
    except Exception:
        return []


def _extract_text_txt(data: bytes) -> str:
    try:
        import chardet

        enc = chardet.detect(data or b""
        ).get("encoding") or "utf-8"
        return data.decode(enc, errors="ignore")
    except Exception:
        try:
            return data.decode("utf-8", errors="ignore")
        except Exception:
            return ""


def build_chunks(text: str, chunk_size: int = 1000, chunk_overlap: int = 150) -> list[str]:
    ts = RecursiveCharacterTextSplitter(chunk_size=chunk_size, chunk_overlap=chunk_overlap)
    return cast(list[str], ts.split_text(text or ""))


def process_source(payload: dict[str, Any]) -> dict[str, Any]:
    empresa_id = payload.get("empresa_id")
    empresa_slug = payload.get("empresa_slug")
    doc_id = payload.get("doc_id")
    mime = payload.get("mime")
    path = payload.get("path")
    presigned_url = payload.get("presigned_url")
    tags = payload.get("tags")
    filename = None
    data = _read_bytes_from_source(path, presigned_url)
    kind = _detect_kind(mime, filename)
    result: dict[str, Any] = {
        "empresa_id": empresa_id,
        "empresa_slug": empresa_slug,
        "doc_id": doc_id,
        "mime": mime,
        "kind": kind,
        "tags": tags,
        "text": None,
        "rows": None,
    }
    if not data:
        return result
    if kind == "pdf":
        result["text"] = _extract_text_pdf(data)
    elif kind == "docx":
        result["text"] = _extract_text_docx(data)
    elif kind == "xlsx":
        result["rows"] = _extract_rows_xlsx(data)
    elif kind == "csv":
        result["rows"] = _extract_rows_csv(data)
    elif kind == "text":
        result["text"] = _extract_text_txt(data)
    else:
        result["text"] = ""
    return result
